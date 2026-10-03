"""Build tiny but valid PDFs and DOCX files in memory so tests don't need fixture files."""
import io
import zlib

from PIL import Image, ImageDraw, ImageFont


def text_image(lines: list[str], size=(900, 260), font_size=40) -> Image.Image:
    """Render black text on white — a stand-in for a scan or screenshot."""
    image = Image.new("RGB", size, "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.load_default(size=font_size)
    for i, line in enumerate(lines):
        draw.text((30, 30 + i * (font_size + 40)), line, fill="black", font=font)
    return image


def image_bytes(image: Image.Image, fmt="PNG") -> bytes:
    buf = io.BytesIO()
    image.save(buf, format=fmt)
    return buf.getvalue()


def make_pdf(pages: list) -> bytes:
    """Return PDF bytes, one page per item.

    Each item is either a string (text; "" = blank page) or a dict:
      {"text": str, "image": PIL.Image, "rect": (x, y, w, h)}  # rect in points, origin bottom-left
    """
    pages = [p if isinstance(p, dict) else {"text": p} for p in pages]
    objects: list[bytes] = [b"", b""]  # 1: catalog, 2: pages (filled in below)
    font_id = None

    def add(body: bytes) -> int:
        objects.append(body)
        return len(objects)

    font_id = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    page_ids = []
    for page in pages:
        text = page.get("text", "")
        ops = []
        xobjects = ""
        if page.get("image") is not None:
            img = page["image"].convert("RGB")
            data = zlib.compress(img.tobytes())
            img_id = add(
                (
                    f"<< /Type /XObject /Subtype /Image /Width {img.width} /Height {img.height} "
                    f"/ColorSpace /DeviceRGB /BitsPerComponent 8 /Filter /FlateDecode /Length {len(data)} >>\nstream\n"
                ).encode()
                + data
                + b"\nendstream"
            )
            x, y, w, h = page.get("rect", (0, 0, 612, 792))
            ops.append(f"q {w} 0 0 {h} {x} {y} cm /Im1 Do Q")
            xobjects = f"/XObject << /Im1 {img_id} 0 R >>"
        if text:
            safe = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            ops.append(f"BT /F1 12 Tf 72 740 Td ({safe}) Tj ET")
        stream = "\n".join(ops).encode()
        content_id = add(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
        page_ids.append(
            add(
                (
                    f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                    f"/Resources << /Font << /F1 {font_id} 0 R >> {xobjects} >> /Contents {content_id} 0 R >>"
                ).encode()
            )
        )

    objects[0] = b"<< /Type /Catalog /Pages 2 0 R >>"
    kids = " ".join(f"{pid} 0 R" for pid in page_ids)
    objects[1] = f"<< /Type /Pages /Kids [{kids}] /Count {len(page_ids)} >>".encode()

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for idx, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{idx} 0 obj\n".encode() + body + b"\nendobj\n"

    xref_at = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_at}\n%%EOF\n".encode()
    return bytes(out)


def make_docx(blocks: list) -> bytes:
    """Build a .docx. Blocks: str (paragraph), list[list[str]] (table), or PIL.Image (inline picture)."""
    from docx import Document
    from docx.shared import Inches

    doc = Document()
    for block in blocks:
        if isinstance(block, str):
            doc.add_paragraph(block)
        elif isinstance(block, Image.Image):
            doc.add_picture(io.BytesIO(image_bytes(block)), width=Inches(5))
        else:
            table = doc.add_table(rows=len(block), cols=len(block[0]))
            for r, row in enumerate(block):
                for c, value in enumerate(row):
                    table.cell(r, c).text = value
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
